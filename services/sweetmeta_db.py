"""Read-only access to the versioned SweetMeta SQLite release."""

from __future__ import annotations

import csv
import hashlib
import os
import sqlite3
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Iterator

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DB_PATH = ROOT / "outputs" / "sweetmeta-reviewed-v107" / "SweetDatabase_v1.0.7_20261005.sqlite"
DEFAULT_SOURCE_CSV = ROOT / "data" / "sweetmeta_reviewed_20261005" / "compound_master_1296.csv"
PUBLIC_LITERATURE_STATUSES = ("accepted", "reviewed")
TIER_NAMES = {
    "R1": "R1_high_readiness",
    "R2": "R2_traceable",
    "R3": "R3_dataset_supported",
    "R4": "R4_high_risk_review",
}


class SweetMetaDatabase:
    """Small, parameterized query layer for the public SweetMeta release."""

    def __init__(
        self,
        path: str | os.PathLike[str] | None = None,
        source_workbook: str | os.PathLike[str] | None = None,
    ) -> None:
        self.path = Path(path or os.getenv("SWEETMETA_DB_PATH", DEFAULT_DB_PATH)).expanduser()
        source_path = Path(
            source_workbook or os.getenv("SWEETMETA_SOURCE_CSV", DEFAULT_SOURCE_CSV)
        ).expanduser()
        self._display_order = self._load_display_order(source_path)

    @staticmethod
    def _load_display_order(source_path: Path) -> dict[str, int]:
        """Freeze public display IDs to the reviewed master row order."""
        if not source_path.is_file():
            return {}
        if source_path.suffix.lower() == ".csv":
            with source_path.open(encoding="utf-8-sig", newline="") as handle:
                return {
                    str(row["compound_id"]): index
                    for index, row in enumerate(csv.DictReader(handle), start=1)
                    if row.get("compound_id")
                }
        # Compatibility for callers that still provide the historical workbook.
        from openpyxl import load_workbook
        workbook = load_workbook(source_path, read_only=True, data_only=True)
        try:
            sheet = workbook["current_sweet_compounds_1296"]
            headers = [str(cell.value or "") for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            column = headers.index("compound_id")
            return {str(row[column]): index for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=1) if row[column]}
        finally:
            workbook.close()

    @property
    def available(self) -> bool:
        return self.path.is_file()

    @contextmanager
    def connection(self) -> Iterator[sqlite3.Connection]:
        if not self.available:
            raise FileNotFoundError(f"SweetMeta SQLite database not found: {self.path}")
        conn = sqlite3.connect(
            f"file:{self.path.resolve()}?mode=ro&immutable=1", uri=True, timeout=5,
        )
        conn.row_factory = sqlite3.Row
        conn.create_function(
            "sweetmeta_display_order",
            1,
            lambda compound_id: self._display_order.get(str(compound_id), 1_000_000),
        )
        conn.create_function(
            "sweetmeta_display_id",
            1,
            lambda compound_id: (
                f"{self._display_order[str(compound_id)]:06d}"
                if str(compound_id) in self._display_order else None
            ),
        )
        try:
            yield conn
        finally:
            conn.close()

    @staticmethod
    def _row(row: sqlite3.Row | None) -> dict[str, Any] | None:
        return dict(row) if row is not None else None

    def _public_where(self) -> tuple[str, tuple[Any, ...]]:
        return (
            "c.compound_id IN (SELECT compound_id FROM release_v1_compound_status)",
            (),
        )

    def integrity(self) -> dict[str, Any]:
        with self.connection() as conn:
            check = conn.execute("PRAGMA integrity_check").fetchone()[0]
            tables = conn.execute("SELECT count(*) FROM sqlite_master WHERE type='table'").fetchone()[0]
            views = conn.execute("SELECT count(*) FROM sqlite_master WHERE type='view'").fetchone()[0]
            compounds = conn.execute("SELECT count(*) FROM compounds").fetchone()[0]
            release = conn.execute("SELECT count(*) FROM release_v1_compound_status").fetchone()[0]
            return {
                "integrity": check,
                "tables": tables,
                "views": views,
                "compounds": compounds,
                "publicRelease": release,
            }

    def metadata(self) -> dict[str, Any]:
        with self.connection() as conn:
            def manifest(table: str) -> dict[str, str]:
                try:
                    return {str(k): str(v) for k, v, *_ in conn.execute(f"SELECT * FROM {table}")}
                except sqlite3.Error:
                    return {}

            counts = {}
            for table in ("literature_records", "literature_mentions", "sweetness_measurements", "food_occurrences", "regulatory_evaluations", "receptor_activities"):
                try:
                    counts[table] = conn.execute(f"SELECT count(*) FROM {table}").fetchone()[0]
                except sqlite3.Error:
                    counts[table] = 0
            rescan_release: dict[str, Any] = {}
            review_release: dict[str, Any] = {}
            reviewed_release = manifest("reviewed_release_metadata")
            try:
                run = conn.execute("SELECT * FROM literature_rescan_runs ORDER BY created_at DESC LIMIT 1").fetchone()
                if run:
                    rescan_release = {**dict(run), "version": "v1.0.6", "databaseFile": self.path.name}
            except sqlite3.Error:
                pass
            try:
                run = conn.execute("SELECT * FROM literature_ai_review_runs ORDER BY created_at DESC LIMIT 1").fetchone()
                if run:
                    review_release = {**dict(run), "version": "v1.0.7", "databaseFile": self.path.name}
            except sqlite3.Error:
                pass
            return {
                "databasePath": str(self.path),
                "databaseSha256": hashlib.sha256(self.path.read_bytes()).hexdigest(),
                "release": manifest("release_v102_manifest"),
                "releaseV1": manifest("release_v1_manifest"),
                "rescanRelease": rescan_release,
                "reviewRelease": review_release,
                "reviewedRelease": {
                    **reviewed_release,
                    "version": reviewed_release.get("version", "v1.0.7"),
                    "databaseFile": self.path.name,
                } if reviewed_release else {},
                "integrity": self.integrity(),
                "counts": counts,
            }

    def stats(self) -> dict[str, Any]:
        with self.connection() as conn:
            tier = conn.execute("""
                SELECT release_evidence_readiness_tier AS tier, count(*) AS count
                FROM release_v1_compound_status GROUP BY tier ORDER BY tier
            """).fetchall()
            gaps = conn.execute("""
                SELECT release_evidence_gap_type AS gap, count(*) AS count
                FROM release_v1_compound_status GROUP BY gap ORDER BY gap
            """).fetchall()
            review = conn.execute("""
                SELECT manual_review_required AS required, count(*) AS count
                FROM release_v1_compound_status GROUP BY required ORDER BY required
            """).fetchall()
            result = self.metadata()
            result.update({
                "publicCompoundCount": conn.execute("SELECT count(*) FROM release_v1_compound_status").fetchone()[0],
                "tierCounts": {row["tier"]: row["count"] for row in tier},
                "gapCounts": {row["gap"]: row["count"] for row in gaps},
                "reviewCounts": {str(row["required"]): row["count"] for row in review},
                "pubchemMatched": conn.execute("SELECT count(*) FROM compound_external_ids WHERE namespace='pubchem'").fetchone()[0],
                "literatureCount": conn.execute("SELECT count(*) FROM literature_records").fetchone()[0],
                "pmidCount": conn.execute("SELECT count(*) FROM literature_records WHERE pmid IS NOT NULL AND trim(pmid) <> ''").fetchone()[0],
            })
            return result

    def statistics(self) -> dict[str, Any]:
        """Return public-release aggregates used by the statistics dashboard."""
        with self.connection() as conn:
            public_total = conn.execute("SELECT count(*) FROM release_v1_compound_status").fetchone()[0]

            def coverage(query: str, params: tuple[Any, ...] = ()) -> tuple[int, int]:
                row = conn.execute(query, params).fetchone()
                return int(row[0] or 0), int(row[1] or 0)

            taste_records, taste_compounds = coverage("""
                SELECT count(*), count(DISTINCT t.compound_id)
                FROM taste_assertions t JOIN release_v1_compound_status r ON r.compound_id=t.compound_id
                WHERE t.assertion_status='accepted'
            """)
            measurement_records, measurement_compounds = coverage("""
                SELECT count(*), count(DISTINCT m.compound_id)
                FROM sweetness_measurements m JOIN release_v1_compound_status r ON r.compound_id=m.compound_id
                WHERE m.log10_relative_sweetness IS NOT NULL
            """)
            literature_records, literature_compounds = coverage("""
                SELECT count(*), count(DISTINCT l.compound_id)
                FROM literature_compound_links l JOIN release_v1_compound_status r ON r.compound_id=l.compound_id
                WHERE l.review_status IN ('accepted','reviewed')
            """)
            food_records, food_compounds = coverage("""
                SELECT count(*), count(DISTINCT f.compound_id)
                FROM food_occurrences f JOIN release_v1_compound_status r ON r.compound_id=f.compound_id
            """)
            regulatory_records, regulatory_compounds = coverage("""
                SELECT count(*), count(DISTINCT s.compound_id)
                FROM regulatory_substances s JOIN release_v1_compound_status r ON r.compound_id=s.compound_id
            """)
            receptor_records, receptor_compounds = coverage("""
                SELECT count(*), count(DISTINCT a.compound_id)
                FROM receptor_activities a JOIN release_v1_compound_status r ON r.compound_id=a.compound_id
            """)
            conflict_records, conflict_compounds = coverage("""
                SELECT count(*), count(DISTINCT d.compound_id)
                FROM data_conflicts d JOIN release_v1_compound_status r ON r.compound_id=d.compound_id
                WHERE d.resolution_status NOT LIKE 'resolved%'
            """)
            quality_issue_compounds = conn.execute("""
                SELECT count(DISTINCT compound_id) FROM (
                    SELECT compound_id FROM release_v1_compound_status WHERE manual_review_required=1
                    UNION
                    SELECT d.compound_id FROM data_conflicts d
                    JOIN release_v1_compound_status r ON r.compound_id=d.compound_id
                    WHERE d.resolution_status NOT LIKE 'resolved%'
                )
            """).fetchone()[0]

            readiness = []
            for row in conn.execute("""
                SELECT release_evidence_readiness_tier AS key, count(*) AS count
                FROM release_v1_compound_status GROUP BY 1 ORDER BY 1
            """):
                readiness.append({
                    "key": str(row["key"]).split("_", 1)[0],
                    "label": str(row["key"]).replace("_", " "),
                    "count": row["count"],
                    "percent": round(row["count"] * 100 / public_total, 2),
                })

            timeline = [dict(row) for row in conn.execute("""
                SELECT (year / 5) * 5 AS yearFrom, (year / 5) * 5 + 4 AS yearTo,
                       coalesce(nullif(trim(sweetener_class), ''), 'other-uncertain') AS family,
                       count(*) AS records
                FROM literature_records WHERE year IS NOT NULL
                GROUP BY 1, 2, 3 ORDER BY 1, 3
            """)]

            sweetness_bins = [
                ("lt-1x", "<1x", None, 0.0),
                ("1-10x", "1-10x", 0.0, 1.0),
                ("10-100x", "10-100x", 1.0, 2.0),
                ("100-1k", "100-1,000x", 2.0, 3.0),
                ("1k-10k", "1,000-10,000x", 3.0, 4.0),
                ("gte-10k", "10,000x+", 4.0, None),
            ]
            sweetness_distribution = []
            for key, label, lower, upper in sweetness_bins:
                filters = ["m.log10_relative_sweetness IS NOT NULL"]
                params: list[Any] = []
                if lower is not None:
                    filters.append("m.log10_relative_sweetness >= ?")
                    params.append(lower)
                if upper is not None:
                    filters.append("m.log10_relative_sweetness < ?")
                    params.append(upper)
                records, compounds = coverage(f"""
                    SELECT count(*), count(DISTINCT m.compound_id)
                    FROM sweetness_measurements m
                    JOIN release_v1_compound_status r ON r.compound_id=m.compound_id
                    WHERE {' AND '.join(filters)}
                """, tuple(params))
                sweetness_distribution.append({
                    "key": key, "label": label, "minLog10": lower, "maxLog10": upper,
                    "records": records, "compounds": compounds,
                })

            domains = [
                ("taste", "Accepted taste evidence", taste_records, taste_compounds),
                ("sweetness", "Quantitative sweetness", measurement_records, measurement_compounds),
                ("literature", "Accepted literature", literature_records, literature_compounds),
                ("food", "Food occurrences", food_records, food_compounds),
                ("regulatory", "Regulatory records", regulatory_records, regulatory_compounds),
                ("receptor", "Receptor activity", receptor_records, receptor_compounds),
                ("conflict", "Unresolved conflicts", conflict_records, conflict_compounds),
            ]
            domain_coverage = [{
                "key": key, "label": label, "records": records, "compounds": compounds,
                "percent": round(compounds * 100 / public_total, 2), "denominator": public_total,
            } for key, label, records, compounds in domains]

            food_groups = [dict(row) for row in conn.execute("""
                SELECT coalesce(nullif(trim(i.food_group), ''), 'Unclassified') AS key,
                       count(*) AS records, count(DISTINCT f.food_id) AS foods,
                       count(DISTINCT f.compound_id) AS compounds
                FROM food_occurrences f JOIN food_items i ON i.food_id=f.food_id
                JOIN release_v1_compound_status r ON r.compound_id=f.compound_id
                GROUP BY 1 ORDER BY records DESC, key
            """)]
            for item in food_groups:
                item["label"] = item["key"]
                item["percent"] = round(item["compounds"] * 100 / public_total, 2)
                item["denominator"] = public_total

            def qc_counts(column: str) -> dict[str, int]:
                values = {str(row["value"] or "unknown"): row["count"] for row in conn.execute(f"""
                    SELECT coalesce(nullif(trim(q.{column}), ''), 'unknown') AS value, count(*) AS count
                    FROM release_v1_compound_status r
                    LEFT JOIN compound_final_qc_v1 q ON q.compound_id=r.compound_id
                    GROUP BY 1
                """)}
                return {"verified": values.get("Y", 0), "attention": values.get("N", 0), "unknown": values.get("unknown", 0)}

            training = {str(row["value"] or "unknown"): row["count"] for row in conn.execute("""
                SELECT coalesce(nullif(trim(q.qc_training_use_status), ''), 'unknown') AS value, count(*) AS count
                FROM release_v1_compound_status r LEFT JOIN compound_final_qc_v1 q ON q.compound_id=r.compound_id
                GROUP BY 1
            """)}
            training_eligible = training.get("eligible", 0)
            review_required = conn.execute("SELECT count(*) FROM release_v1_compound_status WHERE manual_review_required=1").fetchone()[0]
            quality_profile = [
                {"key": "structure", "label": "Structure identity", **qc_counts("structure_verified")},
                {"key": "stereochemistry", "label": "Stereochemistry", **qc_counts("stereochemistry_verified")},
                {"key": "form", "label": "Chemical form", **qc_counts("form_verified")},
                {"key": "training", "label": "Training eligibility", "verified": training_eligible, "attention": public_total - training_eligible, "unknown": 0},
                {"key": "review", "label": "Manual review", "verified": public_total - review_required, "attention": review_required, "unknown": 0},
                {"key": "conflict", "label": "Conflict status", "verified": public_total - conflict_compounds, "attention": conflict_compounds, "unknown": 0},
            ]
            for metric in quality_profile:
                metric["denominator"] = public_total

            release_meta = {str(row[0]): str(row[1]) for row in conn.execute("SELECT key, value FROM reviewed_release_metadata")}
            composition = {
                "level1": [dict(row) for row in conn.execute("""
                    SELECT level1 AS key, level1 AS label, count(*) AS count,
                           round(count(*) * 100.0 / ?, 2) AS percent
                    FROM compound_sweetclass GROUP BY level1 ORDER BY count DESC, level1
                """, (public_total,))],
                "level2": [dict(row) for row in conn.execute("""
                    SELECT sc.level1 AS parent, sc.level2 AS key, sc.level2 AS label,
                           count(*) AS count, totals.parent_total AS parentTotal,
                           round(count(*) * 100.0 / totals.parent_total, 2) AS percentWithinParent,
                           round(count(*) * 100.0 / ?, 2) AS percentOverall
                    FROM compound_sweetclass sc
                    JOIN (
                        SELECT level1, count(*) AS parent_total
                        FROM compound_sweetclass GROUP BY level1
                    ) totals ON totals.level1=sc.level1
                    GROUP BY sc.level1, sc.level2, totals.parent_total
                    ORDER BY sc.level1, count DESC, sc.level2
                """, (public_total,))],
            }
            chemical_space = [dict(row) for row in conn.execute("""
                SELECT p.compound_id AS id, p.method, p.x, p.y, p.fingerprint, p.metric,
                       p.random_seed AS randomSeed, c.preferred_name AS name,
                       c.molecular_formula AS formula, pc.molecular_weight AS molecularWeight,
                       sc.level1 AS sweetClassLevel1, sc.level2 AS sweetClassLevel2,
                       pv.evidence_grade AS evidenceGrade, pv.source_collection AS sourceCollection,
                       pv.source_names AS sourceNames, s.release_evidence_readiness_tier AS releaseTier
                FROM compound_projection p JOIN compounds c ON c.compound_id=p.compound_id
                JOIN compound_sweetclass sc ON sc.compound_id=p.compound_id
                JOIN compound_provenance pv ON pv.compound_id=p.compound_id
                JOIN compound_physicochemical pc ON pc.compound_id=p.compound_id
                JOIN release_v1_compound_status s ON s.compound_id=p.compound_id
                ORDER BY p.method, p.compound_id
            """)]
            evidence = {
                "sourceToGrade": [dict(row) for row in conn.execute("SELECT source_collection AS source, evidence_grade AS target, molecule_count AS count FROM reviewed_stat_source_evidence ORDER BY source, target")],
                "gradeToFamily": [dict(row) for row in conn.execute("SELECT evidence_grade AS source, family_short AS target, molecule_count AS count FROM reviewed_stat_evidence_family ORDER BY source, target")],
                "paths": [dict(row) for row in conn.execute("""
                    SELECT pv.source_collection AS source, pv.evidence_grade AS grade,
                           sc.level1 AS family, count(*) AS count
                    FROM compound_provenance pv
                    JOIN compound_sweetclass sc ON sc.compound_id=pv.compound_id
                    GROUP BY pv.source_collection, pv.evidence_grade, sc.level1
                    ORDER BY pv.source_collection, pv.evidence_grade, sc.level1
                """)],
                "familyGrades": [dict(row) for row in conn.execute("SELECT family_short AS family, evidence_grade AS grade, molecule_count AS count, family_total AS total, within_family_percent AS percent FROM reviewed_stat_evidence_by_family ORDER BY family, grade")],
                "modules": [dict(row) for row in conn.execute("SELECT module, display_metric AS displayMetric, native_grain AS nativeGrain, count, color FROM reviewed_stat_module_counts ORDER BY rowid")],
            }

            def property_distribution(column: str, key: str, label: str,
                                      unit: str, axis_precision: int) -> dict[str, Any]:
                groups = []
                for family_row in conn.execute("SELECT DISTINCT level1 FROM compound_sweetclass ORDER BY level1"):
                    family = family_row[0]
                    values = [float(row[0]) for row in conn.execute(f"""
                        SELECT p.{column} FROM compound_physicochemical p
                        JOIN compound_sweetclass sc ON sc.compound_id=p.compound_id
                        WHERE sc.level1=? ORDER BY p.{column}
                    """, (family,))]
                    n = len(values)
                    quantile = lambda fraction: values[round((n - 1) * fraction)]
                    low, high = values[0], values[-1]
                    width = (high - low) / 20 if high > low else 1.0
                    bins = []
                    for index in range(20):
                        start = low + index * width
                        end = high + 1e-12 if index == 19 else start + width
                        count = sum(1 for value in values if start <= value < end)
                        bins.append({"value": round(start + width / 2, 6), "density": round(count / n, 6)})
                    groups.append({"family": family, "n": n, "min": low, "q1": quantile(.25),
                                   "median": quantile(.5), "q3": quantile(.75), "max": high,
                                   "density": bins})
                return {"key": key, "label": label, "unit": unit,
                        "axisPrecision": axis_precision, "groups": groups}

            properties = [
                property_distribution("molecular_weight", "molecularWeight", "Molecular weight", "Da", 0),
                property_distribution("clogp", "clogp", "cLogP", "", 1),
                property_distribution("tpsa", "tpsa", "TPSA", "A2", 0),
                property_distribution("rotatable_bonds", "rotatableBonds", "Rotatable bonds", "count", 0),
                property_distribution("hba", "hba", "H-bond acceptors", "count", 0),
                property_distribution("hbd", "hbd", "H-bond donors", "count", 0),
                property_distribution("bertz_complexity", "bertzComplexity", "Bertz complexity", "", 0),
                property_distribution("fraction_csp3", "fractionCsp3", "Fraction Csp3", "fraction", 2),
            ]
            coverage_path = Path(__file__).resolve().parents[1] / "data" / "sweetmeta_reviewed_20261005" / "statistics" / "feature_coverage_matrix.csv"
            feature_coverage = []
            if coverage_path.is_file():
                with coverage_path.open(encoding="utf-8-sig", newline="") as handle:
                    for row in csv.DictReader(handle):
                        feature_coverage.append({
                            "resource": row.pop("resource"),
                            "entryScope": row.pop("entry_scope"),
                            "snapshot": row.pop("snapshot"),
                            "features": {key: int(value) for key, value in row.items()},
                        })
            benchmark = {
                "splits": [dict(row) for row in conn.execute("SELECT split, negative_count AS negative, positive_count AS positive, total FROM reviewed_benchmark_splits ORDER BY CASE split WHEN 'train' THEN 1 WHEN 'validation' THEN 2 ELSE 3 END")],
                "models": [dict(row) for row in conn.execute("""
                    SELECT model, display_model AS displayModel, split, seed, n, threshold,
                           roc_auc AS rocAuc, pr_auc AS prAuc,
                           balanced_accuracy AS balancedAccuracy, mcc, brier
                    FROM reviewed_benchmark_metrics ORDER BY roc_auc DESC
                """)],
            }
            rescan = conn.execute("SELECT created_at FROM literature_rescan_runs ORDER BY created_at DESC LIMIT 1").fetchone()
            return {
                "release": {
                    "version": release_meta.get("version", "v1.0.7"), "databaseFile": self.path.name,
                    "sha256": hashlib.sha256(self.path.read_bytes()).hexdigest(),
                    "generatedAt": release_meta.get("created_at") or (rescan[0] if rescan else None),
                    "reviewDate": release_meta.get("review_date"), "publicCompounds": public_total,
                },
                "kpis": [
                    {"key": "compounds", "label": "Public compounds", "value": public_total, "detail": "Formally released compound records"},
                    {"key": "taste", "label": "Taste assertions", "value": taste_records, "detail": f"{taste_compounds:,} public compounds"},
                    {"key": "sweetness", "label": "Sweetness measurements", "value": measurement_records, "detail": f"{measurement_compounds:,} public compounds"},
                    {"key": "literature", "label": "Literature records", "value": conn.execute("SELECT count(*) FROM literature_records").fetchone()[0], "detail": "Curated publication metadata"},
                    {"key": "literature-links", "label": "Literature-linked compounds", "value": literature_compounds, "detail": "Accepted or reviewed links"},
                    {"key": "food", "label": "Food-linked compounds", "value": food_compounds, "detail": f"{food_records:,} occurrence records"},
                    {"key": "regulatory", "label": "Regulatory-linked compounds", "value": regulatory_compounds, "detail": f"{regulatory_records:,} substance links"},
                    {"key": "quality", "label": "Quality review queue", "value": quality_issue_compounds, "detail": "Manual review or unresolved conflict"},
                ],
                "readiness": readiness,
                "literatureTimeline": timeline,
                "sweetnessDistribution": sweetness_distribution,
                "domainCoverage": domain_coverage,
                "foodGroups": food_groups,
                "qualityProfile": quality_profile,
                "composition": composition,
                "chemicalSpace": chemical_space,
                "evidenceCoverage": evidence,
                "properties": properties,
                "featureCoverageComparison": feature_coverage,
                "benchmark": benchmark,
            }

    def list_compounds(self, *, query: str = "", page: int = 1, page_size: int = 25,
                       tier: str = "", review: str = "", domain: str = "",
                       food_group: str = "", sweetness_min_log10: float | None = None,
                       sweetness_max_log10: float | None = None, quality_metric: str = "",
                       quality_state: str = "", sweet_class_l1: str = "",
                       sweet_class_l2: str = "", evidence_grade: str = "",
                       source_collection: str = "",
                       mw_min: float | None = None, mw_max: float | None = None,
                       clogp_min: float | None = None, clogp_max: float | None = None,
                       tpsa_min: float | None = None, tpsa_max: float | None = None) -> dict[str, Any]:
        page = max(1, min(int(page), 100000))
        page_size = max(1, min(int(page_size), 100))
        params: list[Any] = []
        clauses = ["c.compound_id IN (SELECT compound_id FROM release_v1_compound_status)"]
        if query.strip():
            clauses.append("""(
                lower(c.compound_id) LIKE ? OR lower(sweetmeta_display_id(c.compound_id)) LIKE ?
                OR lower(c.preferred_name) LIKE ? OR lower(c.parent_inchikey) LIKE ?
                OR lower(c.molecular_formula) LIKE ?
                OR EXISTS (SELECT 1 FROM compound_aliases a WHERE a.compound_id=c.compound_id AND lower(a.alias) LIKE ?)
                OR EXISTS (
                    SELECT 1 FROM compound_external_ids e
                    WHERE e.compound_id=c.compound_id AND lower(e.external_id) LIKE ?
                )
            )""")
            needle = f"%{query.strip().lower()}%"
            params.extend([needle] * 7)
        if tier and tier != "all":
            normalized_tier = TIER_NAMES.get(tier.upper(), tier)
            clauses.append("s.release_evidence_readiness_tier = ?"); params.append(normalized_tier)
        if review == "required": clauses.append("s.manual_review_required = 1")
        if review == "released": clauses.append("s.manual_review_required = 0")
        if sweet_class_l1:
            clauses.append("sc.level1 = ?"); params.append(sweet_class_l1)
        if sweet_class_l2:
            clauses.append("sc.level2 = ?"); params.append(sweet_class_l2)
        if evidence_grade:
            clauses.append("pv.evidence_grade = ?"); params.append(evidence_grade.upper())
        if source_collection:
            clauses.append("pv.source_collection = ?"); params.append(source_collection)
        for column, lower, upper in (
            ("molecular_weight", mw_min, mw_max), ("clogp", clogp_min, clogp_max),
            ("tpsa", tpsa_min, tpsa_max),
        ):
            if lower is not None:
                clauses.append(f"pc.{column} >= ?"); params.append(lower)
            if upper is not None:
                clauses.append(f"pc.{column} <= ?"); params.append(upper)
        domain_clauses = {
            "taste": "EXISTS (SELECT 1 FROM taste_assertions t WHERE t.compound_id=c.compound_id AND t.assertion_status='accepted')",
            "sweetness": "EXISTS (SELECT 1 FROM sweetness_measurements m WHERE m.compound_id=c.compound_id AND m.log10_relative_sweetness IS NOT NULL)",
            "literature": "EXISTS (SELECT 1 FROM literature_compound_links l WHERE l.compound_id=c.compound_id AND l.review_status IN ('accepted','reviewed'))",
            "food": "EXISTS (SELECT 1 FROM food_occurrences f WHERE f.compound_id=c.compound_id)",
            "regulatory": "EXISTS (SELECT 1 FROM regulatory_substances rs WHERE rs.compound_id=c.compound_id)",
            "receptor": "EXISTS (SELECT 1 FROM receptor_activities ra WHERE ra.compound_id=c.compound_id)",
            "conflict": "EXISTS (SELECT 1 FROM data_conflicts d WHERE d.compound_id=c.compound_id AND d.resolution_status NOT LIKE 'resolved%')",
        }
        if domain in domain_clauses:
            clauses.append(domain_clauses[domain])
        if food_group:
            clauses.append("EXISTS (SELECT 1 FROM food_occurrences f JOIN food_items i ON i.food_id=f.food_id WHERE f.compound_id=c.compound_id AND coalesce(nullif(trim(i.food_group),''),'Unclassified')=?)")
            params.append(food_group)
        if sweetness_min_log10 is not None:
            clauses.append("EXISTS (SELECT 1 FROM sweetness_measurements m WHERE m.compound_id=c.compound_id AND m.log10_relative_sweetness >= ?)")
            params.append(sweetness_min_log10)
        if sweetness_max_log10 is not None:
            clauses.append("EXISTS (SELECT 1 FROM sweetness_measurements m WHERE m.compound_id=c.compound_id AND m.log10_relative_sweetness < ?)")
            params.append(sweetness_max_log10)
        qc_columns = {"structure": "structure_verified", "stereochemistry": "stereochemistry_verified", "form": "form_verified"}
        if quality_metric in qc_columns and quality_state in {"verified", "attention", "unknown"}:
            column = qc_columns[quality_metric]
            if quality_state == "verified": clauses.append(f"q.{column}='Y'")
            elif quality_state == "attention": clauses.append(f"q.{column}='N'")
            else: clauses.append(f"coalesce(trim(q.{column}), '')=''")
        elif quality_metric == "training" and quality_state in {"verified", "attention"}:
            clauses.append("q.qc_training_use_status = 'eligible'" if quality_state == "verified" else "coalesce(q.qc_training_use_status, '') <> 'eligible'")
        elif quality_metric == "review" and quality_state in {"verified", "attention"}:
            clauses.append("s.manual_review_required = 0" if quality_state == "verified" else "s.manual_review_required = 1")
        elif quality_metric == "conflict" and quality_state in {"verified", "attention"}:
            conflict = "EXISTS (SELECT 1 FROM data_conflicts d WHERE d.compound_id=c.compound_id AND d.resolution_status NOT LIKE 'resolved%')"
            clauses.append(f"NOT {conflict}" if quality_state == "verified" else conflict)
        where = " AND ".join(clauses)
        with self.connection() as conn:
            reviewed = bool(conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='compound_sweetclass'").fetchone())
            if reviewed:
                joins = """FROM compounds c JOIN release_v1_compound_status s ON s.compound_id=c.compound_id
                    LEFT JOIN compound_final_qc_v1 q ON q.compound_id=c.compound_id
                    JOIN compound_sweetclass sc ON sc.compound_id=c.compound_id
                    JOIN compound_provenance pv ON pv.compound_id=c.compound_id
                    JOIN compound_physicochemical pc ON pc.compound_id=c.compound_id"""
                reviewed_select = """pc.molecular_weight AS molecularWeight, pc.heavy_atom_count AS heavyAtomCount,
                       pc.clogp, pc.tpsa, pc.rotatable_bonds AS rotatableBonds, pc.hba, pc.hbd,
                       sc.level1 AS sweetClassLevel1, sc.level2 AS sweetClassLevel2,
                       pv.evidence_grade AS evidenceGrade, pv.source_collection AS sourceCollection,
                       pv.source_names AS reviewedSourceNames,"""
            else:
                if any((sweet_class_l1, sweet_class_l2, evidence_grade, source_collection,
                        mw_min, mw_max, clogp_min, clogp_max, tpsa_min, tpsa_max)):
                    return {"items": [], "page": page, "pageSize": page_size, "total": 0, "pageCount": 1}
                joins = """FROM compounds c JOIN release_v1_compound_status s ON s.compound_id=c.compound_id
                    LEFT JOIN compound_final_qc_v1 q ON q.compound_id=c.compound_id"""
                reviewed_select = """c.molecular_weight AS molecularWeight, c.heavy_atom_count AS heavyAtomCount,
                       NULL AS clogp, NULL AS tpsa, NULL AS rotatableBonds, NULL AS hba, NULL AS hbd,
                       NULL AS sweetClassLevel1, NULL AS sweetClassLevel2, NULL AS evidenceGrade,
                       NULL AS sourceCollection, NULL AS reviewedSourceNames,"""
            total = conn.execute(f"SELECT count(*) {joins} WHERE {where}", params).fetchone()[0]
            rows = conn.execute(f"""
                SELECT c.compound_id AS id, sweetmeta_display_id(c.compound_id) AS displayId,
                       'small-molecule' AS entityType, c.preferred_name AS name,
                       c.parent_inchikey AS inchiKey, c.parent_isomeric_smiles AS isomericSmiles,
                       c.parent_canonical_smiles AS canonicalSmiles, c.molecular_formula AS formula,
                       {reviewed_select} c.formal_charge AS formalCharge,
                       s.release_evidence_readiness_tier AS releaseTier,
                       s.release_evidence_gap_type AS releaseGap, s.manual_review_required AS reviewRequired,
                       (
                           SELECT group_concat(external_id, char(31)) FROM (
                               SELECT DISTINCT e.external_id
                               FROM compound_external_ids e
                               WHERE e.compound_id=c.compound_id AND lower(e.namespace)='cas'
                               ORDER BY e.external_id
                           )
                       ) AS casNumbers,
                       (
                           SELECT e.external_id FROM compound_external_ids e
                           WHERE e.compound_id=c.compound_id AND lower(e.namespace)='pubchemcid'
                           ORDER BY e.external_id LIMIT 1
                       ) AS pubchemCid,
                       (
                           SELECT cm.chemical_class FROM coconut_memberships cm
                           WHERE cm.compound_id=c.compound_id AND trim(coalesce(cm.chemical_class, ''))<>''
                           ORDER BY cm.annotation_level DESC, cm.chemical_class LIMIT 1
                       ) AS chemicalClass,
                       (
                           SELECT group_concat(source_name, char(31)) FROM (
                               SELECT DISTINCT coalesce(sr.source_name, records.source_id) AS source_name,
                                      CASE WHEN records.source_id=c.preferred_name_source THEN 0 ELSE 1 END AS source_rank
                               FROM source_records records
                               LEFT JOIN source_registry sr ON sr.source_id=records.source_id
                               WHERE records.compound_id=c.compound_id
                               ORDER BY source_rank, lower(source_name)
                           )
                       ) AS sourceNames
                {joins}
                WHERE {where} ORDER BY sweetmeta_display_order(c.compound_id), c.compound_id
                LIMIT ? OFFSET ?
            """, [*params, page_size, (page - 1) * page_size]).fetchall()
        items = []
        for row in rows:
            item = dict(row)
            item["casNumbers"] = item["casNumbers"].split(chr(31)) if item["casNumbers"] else []
            item["sourceNames"] = item["sourceNames"].split(chr(31)) if item["sourceNames"] else []
            item["reviewedSourceNames"] = item["reviewedSourceNames"].split(";") if item["reviewedSourceNames"] else []
            items.append(item)
        return {"items": items, "page": page, "pageSize": page_size, "total": total, "pageCount": max(1, (total + page_size - 1) // page_size)}

    def compound(self, compound_id: str) -> dict[str, Any] | None:
        with self.connection() as conn:
            row = conn.execute("""
                SELECT c.*, sweetmeta_display_id(c.compound_id) AS display_id,
                       s.release_evidence_readiness_tier AS release_tier,
                       s.release_evidence_gap_type AS release_gap, s.manual_review_required,
                       s.status_note, q.*
                FROM compounds c JOIN release_v1_compound_status s ON s.compound_id=c.compound_id
                LEFT JOIN compound_final_qc_v1 q ON q.compound_id=c.compound_id
                WHERE c.compound_id=?
            """, (compound_id,)).fetchone()
            if row is None: return None
            result = dict(row)
            result["aliases"] = [dict(r) for r in conn.execute("SELECT alias, source_id, source_record_id FROM compound_aliases WHERE compound_id=? ORDER BY alias", (compound_id,))]
            result["externalIds"] = [dict(r) for r in conn.execute("SELECT namespace, external_id, source_id, source_record_id FROM compound_external_ids WHERE compound_id=? ORDER BY namespace, external_id", (compound_id,))]
            classification = conn.execute("SELECT * FROM compound_sweetclass WHERE compound_id=?", (compound_id,)).fetchone()
            provenance = conn.execute("SELECT * FROM compound_provenance WHERE compound_id=?", (compound_id,)).fetchone()
            physicochemical = conn.execute("SELECT * FROM compound_physicochemical WHERE compound_id=?", (compound_id,)).fetchone()
            stereo_review = conn.execute("SELECT * FROM compound_stereo_review WHERE compound_id=?", (compound_id,)).fetchone()
            result["classification"] = dict(classification) if classification else None
            result["provenance"] = dict(provenance) if provenance else None
            result["physicochemical"] = dict(physicochemical) if physicochemical else None
            result["projections"] = [dict(r) for r in conn.execute("SELECT method, x, y, fingerprint, metric, random_seed FROM compound_projection WHERE compound_id=? ORDER BY method", (compound_id,))]
            result["stereoReview"] = dict(stereo_review) if stereo_review else None
            result["tasteAssertions"] = [dict(r) for r in conn.execute("SELECT * FROM taste_assertions WHERE compound_id=? ORDER BY assertion_id", (compound_id,)).fetchall()]
            result["measurements"] = [dict(r) for r in conn.execute("SELECT * FROM sweetness_measurements WHERE compound_id=? ORDER BY rowid", (compound_id,)).fetchall()]
            link_table = conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='literature_compound_links'").fetchone()
            if link_table:
                link_columns = {item[1] for item in conn.execute("PRAGMA table_info(literature_compound_links)")}
                review_method = "lcl.review_method" if "review_method" in link_columns else "NULL"
                review_rule_version = "lcl.review_rule_version" if "review_rule_version" in link_columns else "NULL"
                result["literature"] = [dict(r) for r in conn.execute("""
                    SELECT lr.record_id, lr.title, lr.authors, lr.year, lr.journal, lr.doi, lr.pmid,
                           lcl.relation_type, lcl.matched_text, lcl.evidence_excerpt,
                           lcl.confidence, lcl.review_status, lcl.source_id, lcl.reviewed_at,
                           {review_method} AS review_method,
                           {review_rule_version} AS review_rule_version
                    FROM literature_records lr JOIN literature_compound_links lcl ON lcl.record_id=lr.record_id
                    WHERE lcl.compound_id=? AND lcl.review_status IN ('accepted','reviewed')
                    ORDER BY lr.year DESC, lr.record_id
                """.format(review_method=review_method, review_rule_version=review_rule_version), (compound_id,)).fetchall()]
                result["literatureCandidateCount"] = conn.execute("SELECT count(*) FROM literature_compound_links WHERE compound_id=? AND review_status='candidate'", (compound_id,)).fetchone()[0]
                result["literatureReviewRequiredCount"] = conn.execute("SELECT count(*) FROM literature_compound_links WHERE compound_id=? AND review_status IN ('needs_review','candidate')", (compound_id,)).fetchone()[0]
            else:
                result["literature"] = [dict(r) for r in conn.execute("""
                SELECT lr.record_id, lr.title, lr.authors, lr.year, lr.journal, lr.doi, lr.pmid,
                       lm.matched_text, lm.review_status, lm.confidence
                FROM literature_records lr JOIN literature_mentions lm ON lm.record_id=lr.record_id
                WHERE lm.concept_id IN (SELECT concept_id FROM concept_compound_mappings WHERE compound_id=?)
                   OR lr.source_id IN (SELECT DISTINCT sr.source_id FROM source_records sr WHERE sr.compound_id=?)
                ORDER BY lr.year DESC, lr.record_id
            """, (compound_id, compound_id)).fetchall()]
                result["literatureCandidateCount"] = 0
                result["literatureReviewRequiredCount"] = 0
            result["foodOccurrences"] = [dict(r) for r in conn.execute("""
                SELECT f.*, i.* FROM food_occurrences f
                LEFT JOIN food_items i ON i.food_id=f.food_id
                WHERE f.compound_id=? ORDER BY i.food_group, i.food_name LIMIT 100
            """, (compound_id,)).fetchall()]
            result["regulatory"] = [dict(r) for r in conn.execute("""
                SELECT rs.*, re.* FROM regulatory_substances rs
                LEFT JOIN regulatory_evaluations re ON re.regulatory_substance_id=rs.regulatory_substance_id
                WHERE rs.compound_id=? ORDER BY rs.regulatory_substance_id
            """, (compound_id,)).fetchall()]
            result["receptorActivities"] = [dict(r) for r in conn.execute("SELECT * FROM receptor_activities WHERE compound_id=?", (compound_id,)).fetchall()]
            result["moduleCounts"] = {
                "sweetness": conn.execute("SELECT count(*) FROM sweetness_measurements WHERE compound_id=?", (compound_id,)).fetchone()[0],
                "literature": len(result["literature"]),
                "food": conn.execute("SELECT count(*) FROM food_occurrences WHERE compound_id=?", (compound_id,)).fetchone()[0],
                "receptors": conn.execute("SELECT count(*) FROM receptor_activities WHERE compound_id=?", (compound_id,)).fetchone()[0],
                "regulatory": conn.execute("SELECT count(*) FROM regulatory_substances WHERE compound_id=?", (compound_id,)).fetchone()[0],
                "conflicts": conn.execute("SELECT count(*) FROM data_conflicts WHERE compound_id=?", (compound_id,)).fetchone()[0],
            }
            return result

    def projection_context(self, compound_id: str) -> dict[str, Any] | None:
        """Return only the selected compound's SweetClass projection context."""
        with self.connection() as conn:
            selected = conn.execute("""
                SELECT sc.level1, sc.level2 FROM compound_sweetclass sc
                JOIN release_v1_compound_status s ON s.compound_id=sc.compound_id
                WHERE sc.compound_id=?
            """, (compound_id,)).fetchone()
            if selected is None:
                return None
            points = [dict(row) for row in conn.execute("""
                SELECT p.compound_id AS id, p.method, p.x, p.y,
                       c.preferred_name AS name, sc.level1 AS sweetClassLevel1,
                       sc.level2 AS sweetClassLevel2,
                       CASE WHEN p.compound_id=? THEN 1 ELSE 0 END AS selected
                FROM compound_projection p
                JOIN compounds c ON c.compound_id=p.compound_id
                JOIN compound_sweetclass sc ON sc.compound_id=p.compound_id
                JOIN release_v1_compound_status s ON s.compound_id=p.compound_id
                WHERE sc.level1=?
                ORDER BY p.method, p.compound_id
            """, (compound_id, selected["level1"])).fetchall()]
            return {
                "compoundId": compound_id,
                "family": selected["level1"],
                "subclass": selected["level2"],
                "points": points,
            }

    def literature(self, *, query: str = "", page: int = 1, page_size: int = 25,
                   year: int | None = None, year_from: int | None = None,
                   year_to: int | None = None, pmid_only: bool = False,
                   doi_only: bool = False, review_status: str = "",
                   relation_type: str = "", compound_id: str = "",
                   compound_family: str = "") -> dict[str, Any]:
        page = max(1, int(page)); page_size = max(1, min(int(page_size), 100))
        params: list[Any] = []
        where = "1=1"
        if query.strip():
            needle = f"%{query.strip().lower()}%"
            where = "lower(title) LIKE ? OR lower(coalesce(authors,'')) LIKE ? OR lower(coalesce(journal,'')) LIKE ? OR lower(coalesce(doi,'')) LIKE ? OR lower(coalesce(pmid,'')) LIKE ? OR lower(coalesce(sweetener_class,'')) LIKE ?"
            params.extend([needle] * 6)
        filters = [where]
        has_links = False
        if year is not None:
            filters.append("year = ?"); params.append(year)
        elif year_from is not None:
            filters.append("year >= ?"); params.append(year_from)
        if year_to is not None:
            filters.append("year <= ?"); params.append(year_to)
        if pmid_only:
            filters.append("pmid IS NOT NULL AND trim(pmid) <> ''")
        if doi_only:
            filters.append("doi IS NOT NULL AND trim(doi) <> ''")
        if review_status in {"accepted", "reviewed", "candidate", "needs_review", "rejected"}:
            has_links = True; filters.append("EXISTS (SELECT 1 FROM literature_compound_links l WHERE l.record_id=literature_records.record_id AND l.review_status = ?)"); params.append(review_status)
        elif review_status:
            filters.append("manual_review_required = ?"); params.append(1 if review_status in {"required", "1"} else 0)
        if relation_type:
            has_links = True; filters.append("EXISTS (SELECT 1 FROM literature_compound_links l WHERE l.record_id=literature_records.record_id AND l.relation_type = ? AND l.review_status IN ('accepted','reviewed'))"); params.append(relation_type)
        if compound_id:
            has_links = True; filters.append("EXISTS (SELECT 1 FROM literature_compound_links l WHERE l.record_id=literature_records.record_id AND l.compound_id = ? AND l.review_status IN ('accepted','reviewed'))"); params.append(compound_id)
        if compound_family:
            filters.append("sweetener_class = ?"); params.append(compound_family)
        where = " AND ".join(f"({item})" for item in filters)
        with self.connection() as conn:
            links_available = bool(conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='literature_compound_links'").fetchone())
            quality_available = bool(conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='literature_quality'").fetchone())
            if has_links and not links_available:
                return {"items": [], "page": page, "pageSize": page_size, "pageCount": 1, "total": 0}
            total = conn.execute(f"SELECT count(*) FROM literature_records WHERE {where}", params).fetchone()[0]
            quality_select = "(SELECT quality_status FROM literature_quality q WHERE q.record_id=literature_records.record_id)" if quality_available else "NULL"
            rescan_available = bool(conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='literature_rescan_record_status'").fetchone())
            rescan_select = "(SELECT scan_status FROM literature_rescan_record_status rs WHERE rs.record_id=literature_records.record_id)" if rescan_available else "NULL"
            accepted_select = "(SELECT count(*) FROM literature_compound_links l WHERE l.record_id=literature_records.record_id AND l.review_status IN ('accepted','reviewed'))" if links_available else "0"
            reviewed_select = "(SELECT count(*) FROM literature_compound_links l WHERE l.record_id=literature_records.record_id AND l.review_status='reviewed')" if links_available else "0"
            candidate_select = "(SELECT count(*) FROM literature_compound_links l WHERE l.record_id=literature_records.record_id AND l.review_status IN ('candidate','needs_review'))" if links_available else "0"
            rows = conn.execute(f"""
                SELECT record_id, title, authors, first_author, year, journal, doi, pmid,
                       article_type, study_domain, sweetener_class AS compound_family, research_theme, data_value_tier,
                       language, oa_status, source_url, fulltext_quality_tier,
                       manual_review_required,
                       {quality_select} AS quality_status,
                       {rescan_select} AS rescan_status,
                       {accepted_select} AS accepted_link_count,
                       {reviewed_select} AS reviewed_link_count,
                       {candidate_select} AS candidate_link_count
                FROM literature_records WHERE {where}
                ORDER BY coalesce(year, 0) DESC, record_id LIMIT ? OFFSET ?
            """, [*params, page_size, (page - 1) * page_size]).fetchall()
        return {"items": [dict(row) for row in rows], "page": page, "pageSize": page_size, "pageCount": max(1, (total + page_size - 1) // page_size), "total": total}

    def literature_detail(self, record_id: str) -> dict[str, Any] | None:
        with self.connection() as conn:
            row = conn.execute("SELECT * FROM literature_records WHERE record_id=?", (record_id,)).fetchone()
            if row is None:
                return None
            result = dict(row)
            quality = conn.execute("SELECT * FROM literature_quality WHERE record_id=?", (record_id,)).fetchone()
            result["quality"] = dict(quality) if quality else None
            rescan_available = bool(conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='literature_rescan_record_status'").fetchone())
            rescan = conn.execute("SELECT * FROM literature_rescan_record_status WHERE record_id=?", (record_id,)).fetchone() if rescan_available else None
            result["rescanStatus"] = dict(rescan) if rescan else None
            result["links"] = [dict(link) for link in conn.execute("""
                SELECT l.*, c.preferred_name AS compound_name
                FROM literature_compound_links l JOIN compounds c ON c.compound_id=l.compound_id
                WHERE l.record_id=? AND l.review_status IN ('accepted','reviewed') ORDER BY l.confidence DESC
            """, (record_id,)).fetchall()] if conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='literature_compound_links'").fetchone() else []
            return result
